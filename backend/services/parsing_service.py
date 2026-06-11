import io
import re
from pathlib import Path
from typing import List


def parse_file(content: bytes, filename: str, file_type: str) -> str:
    """Parse uploaded file bytes to Markdown text."""
    if file_type == "pdf":
        return _parse_pdf(content)
    elif file_type == "pptx":
        return _parse_pptx(content)
    elif file_type == "csv":
        return _parse_csv(content)
    elif file_type in ("txt", "md"):
        return content.decode("utf-8", errors="replace")
    else:
        raise ValueError(f"Unsupported file type: {file_type}")


def _parse_pdf(content: bytes) -> str:
    import pdfplumber
    pages = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"## Page {i}\n\n{text.strip()}")
    return "\n\n".join(pages)


def _parse_pptx(content: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(content))
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        if texts:
            slides.append(f"## Slide {i}\n\n" + "\n\n".join(texts))
    return "\n\n".join(slides)


def _parse_csv(content: bytes) -> str:
    import pandas as pd
    try:
        df = pd.read_csv(io.BytesIO(content))
        return df.to_markdown(index=False)
    except Exception:
        return content.decode("utf-8", errors="replace")


def split_into_chunks(markdown: str, pages_per_chunk: int = 5) -> List[str]:
    """Split Markdown into chunks based on page/slide headings."""
    page_pattern = re.compile(r"^## (?:Page|Slide) \d+", re.MULTILINE)
    splits = list(page_pattern.finditer(markdown))

    if not splits:
        # Fallback: split by character count (~3000 chars ≈ ~750 tokens)
        chunk_size = 3000
        return [markdown[i:i + chunk_size] for i in range(0, len(markdown), chunk_size)]

    segments = []
    for idx, match in enumerate(splits):
        start = match.start()
        end = splits[idx + 1].start() if idx + 1 < len(splits) else len(markdown)
        segments.append(markdown[start:end])

    chunks = []
    for i in range(0, len(segments), pages_per_chunk):
        chunk = "\n\n".join(segments[i:i + pages_per_chunk])
        chunks.append(chunk)

    return chunks
