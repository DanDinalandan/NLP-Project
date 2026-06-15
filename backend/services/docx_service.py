import csv
import io
import json
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _add_rich_run(para, text: str):
    """Split text on **bold** / *italic* markers and add styled runs."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if part.startswith("**") and part.endswith("**"):
            para.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*"):
            para.add_run(part[1:-1]).italic = True
        else:
            para.add_run(part)


def markdown_to_docx_bytes(content_markdown: str, title: str = "") -> bytes:
    doc = Document()

    if title:
        h = doc.add_heading(title, level=0)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for raw_line in content_markdown.splitlines():
        line = raw_line.rstrip()

        if line.startswith("#### "):
            doc.add_heading(line[5:], level=4)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            _add_rich_run(p, line[2:])
        elif not line.strip():
            pass  # Word paragraph spacing handles vertical rhythm
        else:
            _add_rich_run(doc.add_paragraph(), line)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def flashcards_to_docx_bytes(flashcards: list, folder_name: str) -> bytes:
    doc = Document()
    doc.add_heading(f"Flashcards — {folder_name}", level=0)

    for i, card in enumerate(flashcards, 1):
        p = doc.add_paragraph()
        p.add_run(f"Q{i}: {card['front']}").bold = True
        doc.add_paragraph(card["back"])

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def mcqs_to_docx_bytes(mcqs: list, folder_name: str) -> bytes:
    doc = Document()
    doc.add_heading(f"MCQ Set — {folder_name}", level=0)
    labels = ["A", "B", "C", "D"]

    for i, q in enumerate(mcqs, 1):
        p = doc.add_paragraph()
        p.add_run(f"{i}. {q['question']}").bold = True

        options = q["options"] if isinstance(q["options"], list) else json.loads(q["options"])
        for j, opt in enumerate(options):
            doc.add_paragraph(f"  {labels[j] if j < 4 else j + 1}. {opt}")

        ans = doc.add_paragraph()
        ans.add_run(f"Answer: {q['correct_answer']}").bold = True
        if q.get("explanation"):
            doc.add_paragraph(f"Explanation: {q['explanation']}")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def mcqs_to_csv_str(mcqs: list) -> str:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(["#", "Question", "Option A", "Option B", "Option C", "Option D", "Correct Answer", "Explanation"])

    for i, q in enumerate(mcqs, 1):
        options = q["options"] if isinstance(q["options"], list) else json.loads(q["options"])
        while len(options) < 4:
            options.append("")
        writer.writerow([
            i,
            q["question"],
            options[0], options[1], options[2], options[3],
            q["correct_answer"],
            q.get("explanation", ""),
        ])

    return out.getvalue()
